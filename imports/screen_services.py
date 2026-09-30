"""IMPORT-001: upload a CSV or Excel file, review every row, then import.

The screen reuses the existing batch pipeline (create -> raw rows -> validate
-> approve -> apply). It adds file reading and three guards the pipeline
lacked (HG-013):
- an opening balance cannot be changed by import once the customer, supplier
  or cashbox has been used (that correction is a dated adjustment, HG-002);
- opening stock is imported once per item and location;
- opening stock dated in closed books is refused.
"""

import csv
import io
import json

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from cashboxes.models import Cashbox, OpeningBalanceTarget
from cashboxes.services import target_has_operational_use
from inventory.models import StockMovement, StockMovementType
from master_data.models import Customer, Item, Location, Supplier

from . import columns
from .apply_services import _date, _decimal, _text, apply_import_batch, get_effective_row_data
from .models import ImportBatch, ImportRowStatus
from .services import add_raw_rows, approve_import_batch, create_import_batch, mark_raw_row_validation, refresh_batch_counters
from .validators import validate_import_batch


#: What the screen offers, in the order a shop should import them. Users are
#: managed from Settings -> Users instead (USERS-001).
SCREEN_TYPES = (
    ("categories", "categories.csv", "التصنيفات", "Categories"),
    ("locations", "locations.csv", "المخازن والفروع", "Locations"),
    ("items", "items.csv", "الأصناف", "Items"),
    ("customers", "customers.csv", "العملاء", "Customers"),
    ("suppliers", "suppliers.csv", "الموردون", "Suppliers"),
    ("cashboxes", "cashboxes.csv", "الخزن", "Cashboxes"),
    ("stock", "opening_stock.csv", "مخزون أول المدة", "Opening stock"),
    ("opening_balances", "opening_balances.csv", "الأرصدة الافتتاحية", "Opening balances"),
)
SCREEN_TYPE_CODES = tuple(row[0] for row in SCREEN_TYPES)
MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024


def _normalise_header(name):
    return str(name or "").strip().lstrip("﻿").strip().lower().replace(" ", "_")


def _clean_rows(rows):
    cleaned = []
    for row in rows:
        data = {_normalise_header(key): ("" if value is None else str(value).strip()) for key, value in row.items() if key not in (None, "")}
        if any(data.values()):
            cleaned.append(data)
    return cleaned


def read_upload(uploaded):
    """Rows as dicts from a .csv (UTF-8 or Windows Arabic) or .xlsx file, keyed by normalised header."""

    headers, rows = read_table(uploaded)
    return _clean_rows(rows)


def read_table(uploaded):
    """(headers as written in the file, rows keyed by those headers). IMPORT-002."""

    name = (uploaded.name or "").lower()
    if uploaded.size and uploaded.size > MAX_BYTES:
        raise ValidationError("file_too_big")
    raw = uploaded.read()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:  # a corrupt or password-protected workbook
            raise ValidationError("file_unreadable") from exc
        values = sheet.iter_rows(values_only=True)
        header = next(values, None) or ()
        keys = [str(cell).replace("\ufeff", "").strip() if cell not in (None, "") else "" for cell in header]
        rows = []
        for record in values:
            rows.append({key: ("" if cell is None else (f"{cell:g}" if isinstance(cell, float) else cell.isoformat()[:10] if hasattr(cell, "isoformat") else cell)) for key, cell in zip(keys, record) if key})
    elif name.endswith(".csv"):
        for encoding in ("utf-8-sig", "cp1256"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValidationError("file_unreadable")
        sample = text[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect=dialect)
        keys = [str(cell).replace("\ufeff", "").strip() for cell in next(reader, [])]
        rows = [dict(zip(keys, record)) for record in reader]
    else:
        raise ValidationError("file_type")
    headers = []
    for key in keys:
        if key and key not in headers:
            headers.append(key)
    rows = [{key: ("" if value is None else str(value).strip()) for key, value in row.items() if key in headers} for row in rows]
    rows = [row for row in rows if any(row.values())]
    if not rows or not headers:
        raise ValidationError("file_empty")
    if len(rows) > MAX_ROWS:
        raise ValidationError("file_too_many_rows")
    return headers, rows


def _opening_balance_guard(target_type, data):
    """The Arabic reason this row may not set an opening balance, or None."""

    if target_type == "opening_balances":
        entity_type = _text(data, "entity_type", "account_type", "target_type").lower()
        code = _text(data, "entity_code", "code")
        lookups = {
            "customer": (Customer, "customer_code", OpeningBalanceTarget.CUSTOMER),
            "supplier": (Supplier, "supplier_code", OpeningBalanceTarget.SUPPLIER),
            "cashbox": (Cashbox, "cashbox_code", OpeningBalanceTarget.CASHBOX),
        }
        if entity_type not in lookups:
            return None
        model, field, target = lookups[entity_type]
    elif target_type in ("customers", "suppliers", "cashboxes"):
        model, field, target = {
            "customers": (Customer, "customer_code", OpeningBalanceTarget.CUSTOMER),
            "suppliers": (Supplier, "supplier_code", OpeningBalanceTarget.SUPPLIER),
            "cashboxes": (Cashbox, "cashbox_code", OpeningBalanceTarget.CASHBOX),
        }[target_type]
        code = _text(data, field, "code")
    else:
        return None
    existing = model.objects.filter(**{field: code}).first()
    if existing is None:
        return None
    try:
        new_balance = _decimal(data, "opening_balance", "balance", default=str(existing.opening_balance))
    except ValidationError:
        return None
    if new_balance != existing.opening_balance and target_has_operational_use(target, existing):
        return "الرصيد الافتتاحي ده اتستخدم في حركات بالفعل؛ عدّله من شاشة «تسوية الرصيد الافتتاحي» مش بالاستيراد."
    return None


def _stock_guard(data, batch, seen):
    from closing.services import ensure_period_is_open

    item = Item.objects.filter(item_code=_text(data, "item_code")).first()
    location = Location.objects.filter(location_code=_text(data, "location_code")).first()
    if item is None or location is None:
        return None
    key = (item.pk, location.pk)
    if key in seen:
        return "نفس الصنف والمخزن متكرر في الملف."
    seen.add(key)
    if StockMovement.objects.filter(item=item, location=location, movement_type=StockMovementType.OPENING_STOCK).exists():
        return "الصنف ده ليه مخزون أول مدة في المخزن ده بالفعل؛ الاستيراد مرة تانية هيكرّره."
    movement_date = _date(data, "movement_date", "stock_date", default=batch.go_live_date or timezone.localdate())
    try:
        ensure_period_is_open(movement_date)
    except ValidationError:
        return "تاريخ مخزون أول المدة جوّه فترة مقفولة."
    return None


@transaction.atomic
def upload_batch(*, target_type, uploaded, user, go_live_date=None):
    """Create a batch from a file, validate every row, apply the HG-013 guards."""

    if target_type not in SCREEN_TYPE_CODES:
        raise ValidationError("file_type")
    headers, rows = read_table(uploaded)
    template = columns.is_template(target_type, headers)
    stamp = timezone.now().strftime("%Y%m%d-%H%M%S-%f")
    batch = create_import_batch(
        batch_code=f"IMP-{target_type}-{stamp}"[:80],
        target_type=target_type,
        source_file_name=(uploaded.name or "")[:255],
        go_live_date=go_live_date,
        user=user,
        notes="" if template else json.dumps({"columns": headers, "mapping": "pending"}, ensure_ascii=False),
    )
    if not template:
        # IMPORT-002: another program's file; its columns are matched on the next screen.
        add_raw_rows(batch.pk, rows)
        return batch
    add_raw_rows(batch.pk, _clean_rows(rows))
    return _check(batch)


def _check(batch):
    """Validate every row, then apply the HG-013 guards."""

    target_type = batch.target_type
    validate_import_batch(batch.pk)
    seen = set()
    for raw_row in batch.raw_rows.order_by("row_number"):
        if raw_row.row_status != ImportRowStatus.VALID:
            continue
        data = get_effective_row_data(raw_row)
        reason = _stock_guard(data, batch, seen) if target_type == "stock" else _opening_balance_guard(target_type, data)
        if reason:
            mark_raw_row_validation(raw_row.pk, is_valid=False, errors=[reason])
    batch.refresh_from_db()
    refresh_batch_counters(batch)
    return batch


@transaction.atomic
def import_batch(batch_id, user):
    """Approve and apply a batch whose rows are all valid. All rows, or none."""

    batch = ImportBatch.objects.select_for_update().get(pk=batch_id)
    approve_import_batch(batch.pk)
    applied = apply_import_batch(batch.pk, user=user)
    if batch.target_type == "stock":
        # Item.average_cost is a display cache (HG-003); refresh it from the new movements.
        from inventory.services import recalculate_item_average_cost

        for item in Item.objects.filter(pk__in={movement.item_id for movement in applied}):
            recalculate_item_average_cost(item)
    batch.refresh_from_db()
    batch.status = "imported"
    batch.full_clean()
    batch.save(update_fields=["status", "updated_at"])
    return len(applied)


def pending_columns(batch):
    """The file's columns while they still wait to be matched, else None."""

    try:
        info = json.loads(batch.notes or "{}")
    except ValueError:
        return None
    return info.get("columns") if isinstance(info, dict) and info.get("mapping") == "pending" else None


def column_samples(batch, headers, count=3):
    rows = [row.raw_data for row in batch.raw_rows.order_by("row_number")[:count]]
    return {header: [str(row.get(header, "")) for row in rows] for header in headers}


def preview(batch, mapping, auto_code=False, count=5):
    rows = [row.raw_data for row in batch.raw_rows.order_by("row_number")[:count]]
    return columns.apply(batch.target_type, rows, mapping, auto_code)


@transaction.atomic
def map_columns(batch_id, mapping, user, auto_code=False):
    """Re-key every row of a pending batch by the confirmed mapping, then check it."""

    batch = ImportBatch.objects.select_for_update().get(pk=batch_id)
    headers = pending_columns(batch)
    if headers is None:
        raise ValidationError("mapping_done")
    allowed = {key for key, *_ in columns.fields(batch.target_type)}
    mapping = {header: (mapping.get(header) or "") for header in headers}
    chosen = [key for key in mapping.values() if key]
    if any(key not in allowed for key in chosen):
        raise ValidationError("mapping_unknown")
    if len(chosen) != len(set(chosen)):
        raise ValidationError("mapping_twice")
    if columns.missing_required(batch.target_type, mapping, auto_code):
        raise ValidationError("mapping_missing")
    raw_rows = list(batch.raw_rows.order_by("row_number"))
    for raw_row, data in zip(raw_rows, columns.apply(batch.target_type, [row.raw_data for row in raw_rows], mapping, auto_code)):
        raw_row.raw_data = data
        raw_row.save(update_fields=["raw_data"])
    batch.notes = json.dumps({"columns": headers, "mapping": mapping, "auto_code": auto_code}, ensure_ascii=False)
    batch.save(update_fields=["notes", "updated_at"])
    return _check(batch)
